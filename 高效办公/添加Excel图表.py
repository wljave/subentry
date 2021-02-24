import os
from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference, BarChart3D

# 设置目标文件夹路径
path = './各部门利润表汇总 - 副本2/'

# 获取文件夹下的所有文件名
file_list = os.listdir(path)
# 遍历文件名列表，取得每一个文件名
for file_name in file_list:
    # 拼接文件路径
    file_path = path + file_name
    print('正在处理：' + file_name)
    # 读取工作簿
    wb = load_workbook(file_path)
    # 读取工作簿中的活跃工作表
    ws = wb.active

    # 实例化 LineChart 类，得到 折线图 - LineChart 对象；柱状图 - BarChart3D 对象
    chart1 = LineChart()
    chart2 = BarChart3D()
    # 引用工作表的部分数据
    data = Reference(worksheet=ws, min_row=3, max_row=9, min_col=1, max_col=5)
    # 添加被引用的数据到 LineChart 对象        chart.add_data(data, from_rows=True, titles_from_data=True)
    chart1.add_data(data, from_rows=True, titles_from_data=True)
    chart2.add_data(data, from_rows=True, titles_from_data=True)
    # chart1.varyColors = True
    # chart2.varyColors = True
    # 添加 LineChart 对象到工作表中，指定折线图的位置
    ws.add_chart(chart1, "C12")
    ws.add_chart(chart2, "C29")

    # 引用工作表的表头数据
    cats = Reference(worksheet=ws, min_row=2, max_row=2, min_col=2, max_col=5)
    # 设置类别轴的标签
    chart1.set_categories(cats)
    chart2.set_categories(cats)
    # 设置 x 轴的标题
    chart1.x_axis.title = "季度"
    chart2.x_axis.title = "季度"
    # 设置 y 轴的标题
    chart1.y_axis.title = "利润"
    chart2.y_axis.title = "利润"
    # 改变线条颜色
    chart1.style = 26
    chart2.style = 26

    # 保存文件
    wb.save('./Styles颜色1-48/' + file_name)
# 在终端提示表格绘图结束
print('恭喜你，工作表中的图绘制成功！')
